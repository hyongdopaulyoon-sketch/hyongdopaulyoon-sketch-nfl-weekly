"""NFL 픽 추적기(xlsx) — 정본은 data/picks.csv, 이 파일은 보기용으로 매번 다시 만든다(손으로 고치지 말 것).
  python tracker_xlsx.py  → data/NFL_픽_추적기.xlsx (시트: 픽로그 · 적중률 · 주별)"""
import csv
import os
from collections import defaultdict

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__)); DATA = os.path.join(os.path.dirname(HERE), "data")
PICKS = os.path.join(DATA, "picks.csv"); OUT = os.path.join(DATA, "NFL_픽_추적기.xlsx")
COLS = [("주차", "week"), ("경기", "game"), ("시장", "market"), ("쪽", "side"), ("라인", "line"), ("배당", "odds"), ("구분", "status"),
        ("등급", "grade"), ("모델값", "model_value"), ("엣지", "edge"), ("승률(모델)", "p_win"), ("결과", "result"), ("스코어", "score"),
        ("유닛", "units"), ("베팅 시각", "placed_at"), ("메모", "note"), ("id", "id")]


def fnum(x):
    try:
        return float(x)
    except (TypeError, ValueError):
        return None


def rec(xs):
    w = sum(1 for x in xs if x["result"] == "W"); l = sum(1 for x in xs if x["result"] == "L"); p = sum(1 for x in xs if x["result"] == "P")
    u = sum(fnum(x["units"]) or 0 for x in xs)
    return [len(xs), w, l, p, round(100 * w / (w + l), 1) if w + l else None, round(u, 2)]


def main():
    rows = list(csv.DictReader(open(PICKS, encoding="utf-8-sig", newline=""))) if os.path.exists(PICKS) else []
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "픽로그"
    head = Font(bold=True); fill = PatternFill("solid", fgColor="DDEBF7")
    ws.append([c for c, _ in COLS])
    for c in ws[1]:
        c.font = head; c.fill = fill
    for r in sorted(rows, key=lambda r: (int(r["week"]), r["game"], r["market"])):
        vals = []
        for _, k in COLS:
            v = r.get(k, "")
            if k in ("week",):
                v = int(v)
            elif k in ("line", "model_value", "edge", "p_win", "units") and fnum(v) is not None:
                v = fnum(v)
            vals.append(v)
        ws.append(vals)
    for i, (c, k) in enumerate(COLS, 1):
        ws.column_dimensions[get_column_letter(i)].width = {"game": 11, "side": 12, "grade": 16, "status": 10, "score": 16, "note": 60, "id": 26, "placed_at": 17}.get(k, 9)
    ws.freeze_panes = "A2"
    # 적중률
    s2 = wb.create_sheet("적중률")
    s2.append(["구분", "픽", "승", "패", "푸시", "적중률%", "유닛", "설명"])
    done = [r for r in rows if r["result"] in ("W", "L", "P")]
    placed = [r for r in done if r["status"] == "placed"]
    pub = [r for r in done if r["grade"].startswith("관심(발행)")]
    lean_ = [r for r in done if r["grade"].startswith("방향(발행)")]
    model = [r for r in done if r["status"] != "placed" and r["grade"].startswith(("관찰", "참고"))]
    s2.append(["실베팅(placed)"] + rec(placed) + ["Paul 이 실제로 건 것 — 손익분기 52.4%"])
    s2.append(["발행 판단(페이퍼)"] + rec(pub) + ["발행 세션 관심·소액 관심(9/29~)"])
    s2.append(["발행 방향(페이퍼)"] + rec(lean_) + ["발행 세션 방향 줄 — 패스 포함 전 경기·스프레드+총점(10/1~)"])
    s2.append(["모델 관찰·참고(페이퍼)"] + rec(model) + ["EPA 모델 방향 — 백테스트 45%, 검증용"])
    tz = [r for r in done if r["grade"] == "티저 다리(관찰)"]
    pr = [fnum(r["odds"]) for r in tz if "표시가" in r["note"] and fnum(r["odds"])]
    avg_c = (sum((-o / (100 - o) if o < 0 else 100 / (o + 100)) for o in pr) / len(pr) * 100) if pr else None
    s2.append(["티저 다리(관찰·페이퍼)"] + rec(tz) + [f"웡 티저 다리(10/1~) — 손익분기 73.9%(−120 2팀 환산) · 60다리 보고 · <72% 폐기"
                                                + (f" · 평균 표시가 {avg_c:.1f}c(n {len(pr)})" if pr else "")])
    wn = [r for r in done if r["grade"] == "바람 언더(관찰)"]
    s2.append(["바람 언더(관찰·페이퍼)"] + rec(wn) + ["실외 · 킥오프 예보 지속 15mph+ 언더(10/1~) — 손익분기 52.4% · 40픽 또는 시즌 종료 때 보고 · <52.4% 폐기"])
    # CLV(2026-10-01): 픽 라인 vs 마감 라인(점, + = 우리가 더 좋은 숫자) — 승패보다 빨리 실력이 보인다. 모델 페이퍼는 리프레시마다 쪽이 갱신돼 ≈0 이 정상
    for lab, xs in (("평균 CLV · 실베팅", [r for r in rows if r["status"] == "placed"]), ("평균 CLV · 발행 판단", [r for r in rows if r["grade"].startswith("관심(발행)")]),
                    ("평균 CLV · 발행 방향", [r for r in rows if r["grade"].startswith("방향(발행)")]),
                    ("평균 CLV · 바람 언더", [r for r in rows if r["grade"] == "바람 언더(관찰)"])):
        cv = [fnum(r.get("clv")) for r in xs if fnum(r.get("clv")) is not None]
        s2.append([lab, len(cv), sum(1 for v in cv if v > 0), sum(1 for v in cv if v < 0), sum(1 for v in cv if v == 0),
                   None, round(sum(cv) / len(cv), 2) if cv else None, "픽/우리 쪽 이동/반대/같음 · 「유닛」 칸 = 평균 CLV(점)"])
    s2.append([])
    s2.append(["실베팅 · 시장별"]); [s2.append([f"  {k}"] + rec([r for r in placed if r["market"] == k]) + [""]) for k in ("spread", "total")]
    s2.append(["실베팅 · 등급별"]); [s2.append([f"  {k}"] + rec([r for r in placed if r["grade"] == k]) + [""]) for k in sorted({r["grade"] for r in placed})]
    s2.append([]); s2.append(["모델 엣지 구간(전체 페이퍼 — 모델 검증)"])
    by = defaultdict(list)
    for r in done:
        e = abs(fnum(r["edge"]) or 0); by["≥5" if e >= 5 else "3.5~5" if e >= 3.5 else "2~3.5" if e >= 2 else "<2"].append(r)
    for k in ("<2", "2~3.5", "3.5~5", "≥5"):
        s2.append([f"  |엣지| {k}"] + rec(by.get(k, [])) + [""])
    for c in s2[1]:
        c.font = head; c.fill = fill
    s2.column_dimensions["A"].width = 30; s2.column_dimensions["H"].width = 40
    # 주별
    s3 = wb.create_sheet("주별")
    s3.append(["주차", "실베팅 픽", "승", "패", "푸시", "적중률%", "유닛", "발행 판단 픽", "승", "패", "푸시", "적중률%", "유닛"])
    for wk in sorted({int(r["week"]) for r in done}):
        s3.append([wk] + rec([r for r in placed if int(r["week"]) == wk]) + rec([r for r in pub if int(r["week"]) == wk]))
    for c in s3[1]:
        c.font = head; c.fill = fill
    wb.save(OUT); print(f"wrote {os.path.relpath(OUT, os.path.dirname(HERE))} — 픽로그 {len(rows)}행 · 채점 {len(done)} · 실베팅 {len(placed)}")


if __name__ == "__main__":
    main()
